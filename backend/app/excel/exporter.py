from pathlib import Path

from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.models.agreement_acceptance import AgreementAcceptance
from app.models.course_registration import CourseRegistration
from app.models.contact_message import ContactMessage
from app.models.internship_application import InternshipApplication
from app.services.google_drive_service import upload_excel_to_google_drive

EXCEL_DIR = Path("app/exports")
EXCEL_DIR.mkdir(parents=True, exist_ok=True)

EXCEL_FILE = EXCEL_DIR / "AxelPath_Management.xlsx"


def export_all_data(db: Session):
    workbook = Workbook()

    # Remove default sheet
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    # ============================================================
    # COURSE REGISTRATIONS (includes agreement + enrollment fields)
    # ============================================================

    course_sheet = workbook.create_sheet("Course Registrations")

    registrations = (
        db.query(CourseRegistration)
        .order_by(CourseRegistration.id.asc())
        .all()
    )

    course_headers = [
        "Registration ID",
        "Full Name",
        "Email",
        "Phone",
        "Referral ID",
        "Course ID",
        "Course Slug",
        "Course Title",
        "Amount (₹)",
        "Agreement Accepted",
        "Agreement Version",
        "Razorpay Order ID",
        "Razorpay Payment ID",
        "UTR",
        "Payment Status",
        "Enrollment Status",
        "Google Drive Folder ID",
        "Google Drive Folder URL",
        "Google Drive Status",
        "Created At",
        "Paid At",
    ]

    course_sheet.append(course_headers)

    for r in registrations:
        course_sheet.append([
            r.registration_id,
            r.full_name,
            r.email,
            r.phone,
            r.referral_id,
            r.course_id,
            r.course_slug,
            r.course_title,
            r.amount,
            "Yes" if r.agreement_accepted else "No",
            r.agreement_version or "",
            r.razorpay_order_id or "",
            r.razorpay_payment_id or "",
            r.utr or "",
            r.payment_status,
            r.enrollment_status,
            r.google_drive_folder_id or "",
            r.google_drive_folder_url or "",
            r.google_drive_status or "",
            str(r.created_at) if r.created_at else "",
            str(r.paid_at) if r.paid_at else "",
        ])

    # ============================================================
    # AGREEMENT ACCEPTANCES (audit trail)
    # ============================================================

    agreement_sheet = workbook.create_sheet("Agreement Acceptances")

    acceptances = (
        db.query(AgreementAcceptance)
        .order_by(AgreementAcceptance.id.asc())
        .all()
    )

    agreement_headers = [
        "ID",
        "Registration ID",
        "Agreement Version",
        "Accepted",
        "Accepted At",
        "IP Address",
        "Created At",
        "Updated At",
    ]

    agreement_sheet.append(agreement_headers)

    for a in acceptances:
        agreement_sheet.append([
            a.id,
            a.registration_id,
            a.agreement_version,
            "Yes" if a.agreement_accepted else "No",
            str(a.accepted_at) if a.accepted_at else "",
            a.ip_address or "",
            str(a.created_at) if a.created_at else "",
            str(a.updated_at) if a.updated_at else "",
        ])

    # ============================================================
    # CONTACT MESSAGES
    # ============================================================

    contact_sheet = workbook.create_sheet("Contact Messages")

    contacts = (
        db.query(ContactMessage)
        .order_by(ContactMessage.id.asc())
        .all()
    )

    contact_columns = [
        column.name for column in ContactMessage.__table__.columns
    ]
    contact_sheet.append(contact_columns)

    for contact in contacts:
        contact_sheet.append([
            getattr(contact, column) for column in contact_columns
        ])

    # ============================================================
    # INTERNSHIP APPLICATIONS
    # ============================================================

    internship_sheet = workbook.create_sheet("Internship Applications")

    internships = (
        db.query(InternshipApplication)
        .order_by(InternshipApplication.id.asc())
        .all()
    )

    internship_columns = [
        column.name for column in InternshipApplication.__table__.columns
    ]
    internship_sheet.append(internship_columns)

    for internship in internships:
        internship_sheet.append([
            getattr(internship, column) for column in internship_columns
        ])

    # ============================================================
    # SAVE + UPLOAD
    # ============================================================

    workbook.save(EXCEL_FILE)
    upload_excel_to_google_drive(EXCEL_FILE)

    return EXCEL_FILE